"""Deterministic tests for schema-guided extraction's local data layer."""
from __future__ import annotations

import io
import json
import tempfile
import unittest
from pathlib import Path

from jsonschema import Draft202012Validator

from app.extraction.api import _export_csv, _export_json, _export_xlsx
from app.extraction.projects import ExtractionProjectStore
from app.extraction.service import (
    _cell_issues,
    _merge_batch_observations,
    _normalize_model_record,
    _reconcile_cross_document_conflicts,
    normalize_fields,
    project_json_schema,
)
from app.store import DocMeta, store
from openpyxl import load_workbook


class SchemaTests(unittest.TestCase):
    def test_schema_supports_nullable_typed_and_nested_fields(self):
        fields = normalize_fields([
            {"id": "company", "name": "Company", "data_type": "string", "required": True, "nullable": False},
            {"id": "revenue", "name": "Revenue", "data_type": "number", "unit": "USD", "allowed_values": ["1", "2"]},
            {"id": "contact", "name": "Contact", "data_type": "object",
             "constraints": {"properties": {"name": {"type": "string"}}, "required": ["name"]}},
        ])
        schema = project_json_schema(fields)
        Draft202012Validator.check_schema(schema)
        Draft202012Validator(schema).validate({"company": "Acme", "revenue": None, "contact": {"name": "A"}})
        self.assertEqual(schema["properties"]["revenue"]["anyOf"][0]["enum"], [1.0, 2.0])

    def test_duplicate_ids_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "unique"):
            normalize_fields([
                {"id": "same", "name": "One"},
                {"id": "same", "name": "Two"},
            ])

    def test_invalid_typed_value_is_reported(self):
        field = normalize_fields([{"id": "year", "name": "Year", "data_type": "integer", "required": True}])[0]
        issues = _cell_issues(field, {"normalized_value": "not a year", "status": "supported", "evidence": [{"page": 1}]})
        self.assertTrue(any("validation failed" in issue.lower() for issue in issues))


class EvidenceTests(unittest.TestCase):
    def test_page_quote_is_verified_and_offsets_are_recorded(self):
        doc_id = "unit-test-doc"
        previous = store.docs.get(doc_id)
        previous_parents = store.parents.get(doc_id)
        store.docs[doc_id] = DocMeta(id=doc_id, filename="annual.pdf", kind="pdf", status="ready", n_pages=1)
        store.parents[doc_id] = {f"{doc_id}:pp1": "Revenue was USD 1,250 for FY2025."}
        try:
            field = normalize_fields([{"id": "revenue", "name": "Revenue", "data_type": "number", "unit": "USD"}])
            row = _normalize_model_record({"record_label": "Acme FY2025", "record_key": "Acme FY2025", "fields": {
                "revenue": {"raw_value": "USD 1,250", "normalized_value": 1250, "status": "supported",
                            "evidence": [{"text": "USD 1,250", "page": 1, "start_char": None, "end_char": None}]},
            }}, doc_id, field, None)
            evidence = row["fields"]["revenue"]["evidence"][0]
            self.assertTrue(evidence["verified"])
            self.assertEqual(evidence["start_char"], 12)
            self.assertEqual(evidence["end_char"], 21)
        finally:
            if previous is None:
                store.docs.pop(doc_id, None)
            else:
                store.docs[doc_id] = previous
            if previous_parents is None:
                store.parents.pop(doc_id, None)
            else:
                store.parents[doc_id] = previous_parents

    def test_conflicting_values_keep_evidence_and_clear_selected_value(self):
        rows = [{"id": "row-1", "record_key": "Acme", "document_id": "doc", "fields": {
            "amount": {"status": "supported", "value": "10", "raw_value": "10", "normalized_value": 10,
                       "evidence": [{"text": "Amount 10", "page": 1}]},
        }}]
        incoming = [{"id": "row-2", "record_key": "Acme", "document_id": "doc", "fields": {
            "amount": {"status": "supported", "value": "12", "raw_value": "12", "normalized_value": 12,
                       "evidence": [{"text": "Amount 12", "page": 2}]},
        }}]
        merged = _merge_batch_observations(rows, incoming)
        cell = merged[0]["fields"]["amount"]
        self.assertEqual(cell["status"], "conflicting")
        self.assertIsNone(cell["normalized_value"])
        self.assertEqual(len(cell["evidence"]), 2)

    def test_conflicts_across_documents_are_flagged_with_both_sources(self):
        rows = [
            {"id": "a", "record_key": "Acme FY2025", "schema_version": 1, "document_id": "doc-a", "filename": "a.pdf", "fields": {
                "revenue": {"status": "supported", "value": "10", "raw_value": "10", "normalized_value": 10,
                            "review_status": "unreviewed", "validation_issues": [], "evidence": [{"text": "Revenue 10", "page": 2}]},
            }},
            {"id": "b", "record_key": "Acme FY2025", "schema_version": 1, "document_id": "doc-b", "filename": "b.pdf", "fields": {
                "revenue": {"status": "supported", "value": "12", "raw_value": "12", "normalized_value": 12,
                            "review_status": "unreviewed", "validation_issues": [], "evidence": [{"text": "Revenue 12", "page": 3}]},
            }},
        ]
        _reconcile_cross_document_conflicts(rows)
        self.assertEqual([row["fields"]["revenue"]["status"] for row in rows], ["conflicting", "conflicting"])
        self.assertEqual(len(rows[0]["fields"]["revenue"]["evidence"]), 2)


class PersistenceAndExportTests(unittest.TestCase):
    def test_projects_persist_and_mark_interrupted_jobs(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first = ExtractionProjectStore(root)
            project = first.create("Reports", "Extract revenue")
            first.update(project["id"], lambda p: p["job"].update({"id": "job-1", "status": "running"}))
            restarted = ExtractionProjectStore(root)
            restored = restarted.get(project["id"])
            self.assertEqual(restored["job"]["status"], "error")
            self.assertIn("restarted", restored["job"]["error"].lower())

    def test_json_csv_and_excel_exports_contain_real_rows_and_evidence(self):
        project = {
            "id": "project-1", "name": "Finance", "objective": "Extract revenue",
            "schema": {"version": 1, "fields": [{"id": "revenue", "name": "Revenue", "data_type": "number"}]},
            "observations": [{"id": "row-1", "label": "Acme FY2025", "observation_unit": "company year",
                "document_id": "doc-1", "filename": "report.pdf", "fields": {"revenue": {
                    "value": "=1+1", "normalized_value": 2, "status": "supported", "review_status": "unreviewed",
                    "evidence": [{"text": "Revenue was 2", "page": 3, "start_char": 0, "end_char": 13}],
                    "validation_issues": [],
                }}}],
        }
        encoded = json.loads(_export_json(project))
        self.assertEqual(encoded["observations"][0]["fields"]["revenue"]["normalized_value"], 2)
        csv_text = _export_csv(project).decode("utf-8-sig")
        self.assertIn("'" + "=1+1", csv_text)
        workbook = load_workbook(io.BytesIO(_export_xlsx(project)), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Results", "Evidence", "Validation"])
        self.assertEqual(workbook["Results"][2][6].value, "'=1+1")
        self.assertEqual(workbook["Evidence"][2][10].value, 3)


if __name__ == "__main__":
    unittest.main()
