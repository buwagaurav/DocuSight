"""HTTP API for persistent schema-guided extraction projects."""
from __future__ import annotations

import csv
import io
import json
import re
import time
from dataclasses import asdict
from typing import Any

import pymupdf
from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import Response
from jsonschema import Draft202012Validator
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field

from .. import pipeline
from ..ingest.pdf import _MUPDF_LOCK
from ..retrieval.colpali import colpali
from ..retrieval.hybrid import retriever
from ..config import EXTRACTION_CLOUD_ENABLED
from ..llm.client import llm_available, model_label
from ..store import store
from .projects import projects
from .service import (
    discover_schema,
    estimate_project,
    normalize_fields,
    project_json_schema,
    cancel_job,
    _reconcile_cross_document_conflicts,
    require_cloud_consent,
    resume_job,
    start_job,
    supports_vision,
)

router = APIRouter(prefix="/api/extraction", tags=["schema-guided extraction"])


def _tag_project_documents() -> None:
    """Projects created before documents carried a project_id: tag their uploads on startup."""
    for project in projects.list():
        for doc_id in project.get("documents", []):
            meta = store.docs.get(doc_id)
            if meta and meta.project_id is None:
                store.update(doc_id, project_id=project["id"])


_tag_project_documents()
MAX_UPLOAD = 200 * 1024 * 1024


class ProjectCreate(BaseModel):
    name: str = "Untitled extraction"
    objective: str = ""


class ProjectPatch(BaseModel):
    name: str | None = None
    objective: str | None = None
    observation_unit: str | None = None


class DiscoverRequest(BaseModel):
    objective: str | None = None
    observation_unit: str | None = None
    cloud_consent: bool = False


class SchemaUpdate(BaseModel):
    objective: str | None = None
    observation_unit: str
    fields: list[dict[str, Any]]


class SchemaReuse(BaseModel):
    source_project_id: str


class JobRequest(BaseModel):
    cloud_consent: bool = False
    document_ids: list[str] = Field(default_factory=list)
    field_ids: list[str] = Field(default_factory=list)
    observation_ids: list[str] = Field(default_factory=list)


class ConsentRequest(BaseModel):
    cloud_consent: bool = False


class CellCorrection(BaseModel):
    value: str
    normalized_value: Any = None
    status: str = "supported"


def _project(project_id: str) -> dict[str, Any]:
    value = projects.get(project_id)
    if value is None:
        raise HTTPException(404, "Extraction project not found")
    return value


def _summary(project: dict[str, Any]) -> dict[str, Any]:
    return {k: v for k, v in project.items() if k != "observations"} | {"observation_count": len(project.get("observations", []))}


def _ensure_idle(project: dict[str, Any]) -> None:
    if project.get("job", {}).get("status") in {"queued", "running"}:
        raise HTTPException(409, "Wait for the current extraction job to finish before changing the project")


def _cloud(consent: bool) -> None:
    try:
        require_cloud_consent(consent)
    except ValueError as exc:
        status = 503 if "credentials" in str(exc).lower() else 403
        raise HTTPException(status, str(exc)) from exc


@router.get("/capabilities")
def capabilities() -> dict[str, Any]:
    return {"cloud_enabled": EXTRACTION_CLOUD_ENABLED, "llm_available": llm_available(),
            "provider": model_label(), "local_generation_available": False,
            "vision_available": supports_vision(), "cloud_consent_required": True, "cost_pricing_available": False}


@router.post("/projects")
def create_project(body: ProjectCreate) -> dict[str, Any]:
    return projects.create(body.name, body.objective)


@router.get("/projects")
def list_projects() -> list[dict[str, Any]]:
    rows = []
    for project in projects.list():
        rows.append({k: v for k, v in project.items() if k != "observations"} | {
            "observation_count": len(project.get("observations", [])),
            "field_count": len(project.get("schema", {}).get("fields", [])),
        })
    return rows


@router.get("/projects/{project_id}")
def get_project(project_id: str) -> dict[str, Any]:
    return _summary(_project(project_id))


@router.patch("/projects/{project_id}")
def patch_project(project_id: str, body: ProjectPatch) -> dict[str, Any]:
    _ensure_idle(_project(project_id))
    changes = body.model_dump(exclude_unset=True)
    def mutate(project):
        for key, value in changes.items():
            if value is not None:
                project[key] = value.strip()
    return _summary(projects.update(project_id, mutate) or {})


@router.delete("/projects/{project_id}")
def delete_project(project_id: str) -> dict[str, bool]:
    project = _project(project_id)
    if project.get("job", {}).get("status") in {"queued", "running"}:
        raise HTTPException(409, "Wait for the current extraction job to finish before deleting the project")
    projects.delete(project_id)
    for doc_id in project.get("documents", []):  # project uploads are project-scoped: remove them too
        meta = store.docs.get(doc_id)
        if meta and meta.project_id == project_id:
            store.delete(doc_id)
            colpali.forget(doc_id)
    retriever.invalidate()
    return {"ok": True}


@router.post("/projects/{project_id}/documents")
async def upload_project_documents(project_id: str, files: list[UploadFile] = File(...)) -> list[dict[str, Any]]:
    _ensure_idle(_project(project_id))
    pending: list[tuple[str, bytes]] = []
    for upload in files:
        filename = (upload.filename or "upload.pdf").replace("\\", "/").rsplit("/", 1)[-1]
        if filename.lower().endswith(".pdf") is False:
            raise HTTPException(400, f"{filename}: extraction projects accept PDF files only")
        data = await upload.read()
        if not data:
            raise HTTPException(400, f"{filename} is empty")
        if len(data) > MAX_UPLOAD:
            raise HTTPException(413, f"{filename} is larger than 200 MB")
        try:
            with _MUPDF_LOCK:  # MuPDF is not thread-safe; ingestion may be parsing another PDF right now
                pdf = pymupdf.open(stream=data, filetype="pdf")
                pages = pdf.page_count
                pdf.close()
            if pages < 1:
                raise ValueError("PDF contains no pages")
        except Exception as exc:
            raise HTTPException(400, f"{filename} is not a readable PDF: {exc}") from exc
        pending.append((filename, data))
    created = []
    for filename, data in pending:
        meta = store.create(filename, "pdf", data, project_id=project_id)
        created.append(asdict(meta))
    if created:
        ids = [doc["id"] for doc in created]
        projects.update(project_id, lambda p: p["documents"].extend(i for i in ids if i not in p["documents"]))
        for doc_id in ids:
            # Keep project upload local until the user explicitly approves schema
            # discovery or extraction; skip the RAG app's automatic LLM summary.
            pipeline.submit(doc_id, auto_extract=False, contextual_retrieval=False)
    return created


@router.get("/projects/{project_id}/documents")
def list_project_documents(project_id: str) -> list[dict[str, Any]]:
    project = _project(project_id)
    return [asdict(store.docs[doc_id]) for doc_id in project.get("documents", []) if doc_id in store.docs]


@router.post("/projects/{project_id}/discover")
def discover(project_id: str, body: DiscoverRequest) -> dict[str, Any]:
    project = _project(project_id)
    _ensure_idle(project)
    _cloud(body.cloud_consent)
    objective = (body.objective if body.objective is not None else project.get("objective", "")).strip()
    if not objective:
        raise HTTPException(400, "Enter an extraction objective before discovering a schema")
    doc_ids = [i for i in project.get("documents", []) if i in store.docs and store.docs[i].status == "ready"]
    if not doc_ids:
        raise HTTPException(409, "Upload PDFs and wait for document processing to finish before schema discovery")
    if len(doc_ids) != len(project.get("documents", [])):
        raise HTTPException(409, "Wait until every PDF in this project is ready before schema discovery")
    projects.update(project_id, lambda p: p.update({"objective": objective, "cloud_consent_at": time.time()}))
    try:
        proposed = discover_schema(objective, body.observation_unit or project.get("observation_unit", ""), doc_ids[:3])
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(502, f"Schema discovery failed: {type(exc).__name__}: {exc}") from exc
    projects.update(project_id, lambda p: p.update({"observation_unit": proposed["observation_unit"],
                                                    "discovery_reasoning": proposed["reasoning"],
                                                    "proposed_schema": {"fields": proposed["fields"]}}))
    return proposed


@router.put("/projects/{project_id}/schema")
def update_schema(project_id: str, body: SchemaUpdate) -> dict[str, Any]:
    project = _project(project_id)
    _ensure_idle(project)
    try:
        fields = normalize_fields(body.fields)
        json_schema = project_json_schema(fields)
        Draft202012Validator.check_schema(json_schema)
    except Exception as exc:
        raise HTTPException(400, f"Invalid schema: {exc}") from exc
    version = int(project.get("schema_version", 0)) + 1
    snapshot = {"version": version, "observation_unit": body.observation_unit.strip(),
                "fields": fields, "json_schema": json_schema, "saved_at": time.time()}
    def mutate(p):
        if p.get("schema_version", 0) > 0:
            p.setdefault("schema_versions", []).append(p["schema"])
        p["schema"] = snapshot
        p["schema_version"] = version
        p["observation_unit"] = body.observation_unit.strip()
        if body.objective is not None:
            p["objective"] = body.objective.strip()
        p["status"] = "schema_ready"
    return _summary(projects.update(project_id, mutate) or {})


@router.post("/projects/{project_id}/reuse-schema")
def reuse_schema(project_id: str, body: SchemaReuse) -> dict[str, Any]:
    target = _project(project_id)
    _ensure_idle(target)
    source = _project(body.source_project_id)
    if target["id"] == source["id"]:
        raise HTTPException(400, "Choose a different project as the schema source")
    if source.get("schema_version", 0) < 1 or not source.get("schema", {}).get("fields"):
        raise HTTPException(409, "The source project has no saved schema to reuse")
    fields = normalize_fields(source["schema"]["fields"])
    json_schema = project_json_schema(fields)
    version = int(target.get("schema_version", 0)) + 1
    snapshot = {"version": version, "observation_unit": source.get("observation_unit", ""),
                "fields": fields, "json_schema": json_schema, "saved_at": time.time(),
                "copied_from_project_id": source["id"]}
    def mutate(p):
        if p.get("schema_version", 0) > 0:
            p.setdefault("schema_versions", []).append(p["schema"])
        p["schema"] = snapshot
        p["schema_version"] = version
        p["observation_unit"] = source.get("observation_unit", "")
        p["status"] = "schema_ready"
    return _summary(projects.update(project_id, mutate) or {})


@router.get("/projects/{project_id}/estimate")
def estimate(project_id: str) -> dict[str, Any]:
    return estimate_project(_project(project_id))


@router.post("/projects/{project_id}/extract")
def extract(project_id: str, body: JobRequest) -> dict[str, Any]:
    project = _project(project_id)
    _cloud(body.cloud_consent)
    if not project.get("schema", {}).get("fields"):
        raise HTTPException(409, "Review and save a schema before starting extraction")
    projects.update(project_id, lambda p: p.update({"cloud_consent_at": time.time()}))
    try:
        return _summary(start_job(project_id, document_ids=body.document_ids or None))
    except KeyError as exc:
        raise HTTPException(404, str(exc)) from exc
    except ValueError as exc:
        status = 409 if "already running" in str(exc) else 400
        raise HTTPException(status, str(exc)) from exc


@router.post("/projects/{project_id}/rerun")
def rerun(project_id: str, body: JobRequest) -> dict[str, Any]:
    project = _project(project_id)
    _cloud(body.cloud_consent)
    if not project.get("schema", {}).get("fields"):
        raise HTTPException(409, "Save a schema before rerunning extraction")
    if body.field_ids and body.observation_ids and not body.document_ids:
        known = {row["id"] for row in project.get("observations", [])}
        if not set(body.observation_ids).issubset(known):
            raise HTTPException(400, "Unknown observation ID")
    projects.update(project_id, lambda p: p.update({"cloud_consent_at": time.time()}))
    try:
        return _summary(start_job(project_id, document_ids=body.document_ids or None,
                                  field_ids=body.field_ids or None, observation_ids=body.observation_ids or None))
    except KeyError as exc:
        raise HTTPException(404, str(exc)) from exc
    except ValueError as exc:
        status = 409 if "already running" in str(exc) else 400
        raise HTTPException(status, str(exc)) from exc


@router.post("/projects/{project_id}/resume")
def resume(project_id: str, body: ConsentRequest) -> dict[str, Any]:
    _cloud(body.cloud_consent)
    _project(project_id)
    projects.update(project_id, lambda p: p.update({"cloud_consent_at": time.time()}))
    try:
        return _summary(resume_job(project_id))
    except KeyError as exc:
        raise HTTPException(404, str(exc)) from exc
    except ValueError as exc:
        status = 409 if "already running" in str(exc) else 400
        raise HTTPException(status, str(exc)) from exc


@router.post("/projects/{project_id}/cancel")
def cancel(project_id: str) -> dict[str, Any]:
    _project(project_id)
    try:
        return _summary(cancel_job(project_id))
    except ValueError as exc:
        raise HTTPException(409, str(exc)) from exc


@router.get("/projects/{project_id}/results")
def results(project_id: str, limit: int = 100, offset: int = 0, q: str = "", sort_by: str = "document") -> dict[str, Any]:
    project = _project(project_id)
    all_rows = project.get("observations", [])
    rows = all_rows
    query = q.strip().casefold()
    if query:
        rows = [row for row in rows if query in (str(row.get("label", "")) + " " + str(row.get("filename", "")) + " " +
                " ".join(str(cell.get("value", "")) for cell in row.get("fields", {}).values())).casefold()]
    if sort_by == "observation":
        rows = sorted(rows, key=lambda row: (str(row.get("label", "")).casefold(), str(row.get("filename", "")).casefold()))
    else:
        rows = sorted(rows, key=lambda row: (str(row.get("filename", "")).casefold(), str(row.get("label", "")).casefold()))
    safe_limit = min(max(limit, 1), 1000)
    safe_offset = max(offset, 0)
    return {"total": len(rows), "all_total": len(all_rows), "offset": safe_offset, "limit": safe_limit,
            "observations": rows[safe_offset:safe_offset + safe_limit]}


@router.patch("/projects/{project_id}/results/{observation_id}/fields/{field_id}")
def correct_cell(project_id: str, observation_id: str, field_id: str, body: CellCorrection) -> dict[str, Any]:
    project = _project(project_id)
    _ensure_idle(project)
    field = next((f for f in project.get("schema", {}).get("fields", []) if f["id"] == field_id), None)
    if field is None:
        raise HTTPException(404, "Schema field not found")
    if body.status not in {"supported", "missing", "illegible", "not_applicable", "conflicting", "unsupported"}:
        raise HTTPException(400, "Unsupported cell status")
    result: dict[str, Any] = {}
    def mutate(p):
        row = next((r for r in p.get("observations", []) if r["id"] == observation_id), None)
        if row is None:
            raise HTTPException(404, "Observation not found")
        cell = row.setdefault("fields", {}).setdefault(field_id, {"evidence": [], "validation_issues": []})
        cell.pop("cross_document_original", None)
        if cell.get("original_model_value") is None:
            cell["original_model_value"] = {"value": cell.get("value"), "normalized_value": cell.get("normalized_value"),
                                            "status": cell.get("status")}
        cell.update({"value": body.value, "raw_value": body.value, "normalized_value": body.normalized_value,
                     "status": body.status, "review_status": "corrected", "corrected_at": time.time()})
        cell["validation_issues"] = []
        if field.get("required") and body.normalized_value in (None, "", []):
            cell["validation_issues"].append("Required field is missing.")
        if body.status == "supported" and not cell.get("evidence"):
            cell["validation_issues"].append("No source evidence was returned.")
        try:
            Draft202012Validator({"type": "object", "properties": {field_id: project_json_schema([field])["properties"][field_id]},
                                  "required": [field_id], "additionalProperties": False}).validate({field_id: body.normalized_value})
        except Exception as exc:
            cell["validation_issues"].append(f"Type or constraint validation failed: {exc.message if hasattr(exc, 'message') else exc}")
        _reconcile_cross_document_conflicts(p.get("observations", []))
        result.update(cell)
    updated = projects.update(project_id, mutate)
    if updated is None:
        raise HTTPException(404, "Extraction project not found")
    return result


def _json_value(value: Any) -> str:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return "" if value is None else str(value)


def _spreadsheet_safe(value: Any) -> Any:
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _export_json(project: dict[str, Any]) -> bytes:
    payload = {"project": {k: v for k, v in project.items() if k != "observations"},
               "schema": project.get("schema"), "observations": project.get("observations", []),
               "validation_issues": [
                   {"observation_id": row["id"], "field_id": fid, "issues": cell.get("validation_issues", [])}
                   for row in project.get("observations", []) for fid, cell in row.get("fields", {}).items()
                   if cell.get("validation_issues")
               ]}
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def _export_csv(project: dict[str, Any]) -> bytes:
    output = io.StringIO(newline="")
    fields = project.get("schema", {}).get("fields", [])
    headers = ["observation_id", "observation_unit", "observation_label", "document_id", "filename"]
    field_names: dict[str, str] = {}
    for field in fields:
        name = str(_spreadsheet_safe(field["name"]))
        field_names[field["id"]] = name
        headers += [name, f"{name}__normalized", f"{name}__status", f"{name}__evidence", f"{name}__validation"]
    headers = [_spreadsheet_safe(name) for name in headers]
    writer = csv.DictWriter(output, fieldnames=headers)
    writer.writeheader()
    for row in project.get("observations", []):
        result = {"observation_id": _spreadsheet_safe(row.get("id")), "observation_unit": _spreadsheet_safe(row.get("observation_unit")),
                  "observation_label": _spreadsheet_safe(row.get("label")), "document_id": _spreadsheet_safe(row.get("document_id")),
                  "filename": _spreadsheet_safe(row.get("filename"))}
        for field in fields:
            cell = row.get("fields", {}).get(field["id"], {})
            name = field_names[field["id"]]
            result[name] = _spreadsheet_safe(cell.get("value", ""))
            result[f"{name}__normalized"] = _spreadsheet_safe(_json_value(cell.get("normalized_value")))
            result[f"{name}__status"] = cell.get("status", "missing")
            result[f"{name}__evidence"] = _spreadsheet_safe(_json_value(cell.get("evidence", [])))
            result[f"{name}__validation"] = _spreadsheet_safe(_json_value(cell.get("validation_issues", [])))
        writer.writerow(result)
    return output.getvalue().encode("utf-8-sig")


def _export_xlsx(project: dict[str, Any]) -> bytes:
    workbook = Workbook()
    results_sheet = workbook.active
    results_sheet.title = "Results"
    fields = project.get("schema", {}).get("fields", [])
    headers = ["Observation ID", "Observation unit", "Observation", "Document ID", "Filename"]
    for field in fields:
        headers.extend([field["name"], f"{field['name']} · normalized", f"{field['name']} · status"])
    results_sheet.append([_spreadsheet_safe(value) for value in headers])
    for row in project.get("observations", []):
        values: list[Any] = [row.get("id"), row.get("observation_unit"), row.get("label"), row.get("document_id"), row.get("filename")]
        for field in fields:
            cell = row.get("fields", {}).get(field["id"], {})
            normalized = cell.get("normalized_value")
            if isinstance(normalized, (dict, list)):
                normalized = json.dumps(normalized, ensure_ascii=False)
            values.extend([cell.get("value", ""), normalized, cell.get("status", "missing")])
        results_sheet.append([_spreadsheet_safe(value) for value in values])

    evidence = workbook.create_sheet("Evidence")
    evidence.append([_spreadsheet_safe(value) for value in ["Observation ID", "Observation", "Observation document", "Evidence document ID", "Evidence filename", "Field", "Value", "Status", "Evidence text", "Page", "Start", "End", "Review"]])
    issues = workbook.create_sheet("Validation")
    issues.append([_spreadsheet_safe(value) for value in ["Observation ID", "Observation", "Document", "Field", "Issue", "Review"]])
    for row in project.get("observations", []):
        for field in fields:
            cell = row.get("fields", {}).get(field["id"], {})
            for source in cell.get("evidence", []):
                evidence.append([_spreadsheet_safe(v) for v in [row.get("id"), row.get("label"), row.get("filename"),
                                 source.get("document_id", row.get("document_id")), source.get("filename", row.get("filename")),
                                 field["name"], cell.get("value"), cell.get("status"), source.get("text"), source.get("page"),
                                 source.get("start_char"), source.get("end_char"), cell.get("review_status")]])
            for issue in cell.get("validation_issues", []):
                issues.append([_spreadsheet_safe(v) for v in [row.get("id"), row.get("label"), row.get("filename"), field["name"], issue, cell.get("review_status")]])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1D4ED8")
        for column in sheet.columns:
            letter = column[0].column_letter
            width = min(52, max(12, max((len(str(c.value or "")) for c in column[:50]), default=10) + 2))
            sheet.column_dimensions[letter].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@router.get("/projects/{project_id}/export/{format}")
def export(project_id: str, format: str) -> Response:
    project = _project(project_id)
    if format == "json":
        data, media, suffix = _export_json(project), "application/json", "json"
    elif format == "csv":
        data, media, suffix = _export_csv(project), "text/csv; charset=utf-8", "csv"
    elif format in {"xlsx", "excel"}:
        data, media, suffix = _export_xlsx(project), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    else:
        raise HTTPException(400, "Export format must be json, csv, or xlsx")
    safe_name = re.sub(r"[^A-Za-z0-9_.-]+", "_", project.get("name", "extraction")).strip("._") or "extraction"
    filename = f"{safe_name}.{suffix}"
    return Response(content=data, media_type=media, headers={"Content-Disposition": f'attachment; filename="{filename}"'})
