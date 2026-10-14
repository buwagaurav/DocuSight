"""Excel / CSV ingestion for messy, real-world workbooks.

Handles:
  * merged cells (value propagated to the whole merged range)
  * several tables per sheet (split on blank rows AND blank columns)
  * title / caption rows above a table, notes below it
  * multi-row headers ("2024" over "Q1 Q2 Q3" -> "2024 / Q1")
  * form-style sheets (label | value pairs, e.g. invoices built in Excel)
  * accounting formats ($1,234 / (300) / 12%) -> numeric
  * formulas (cached values are used; formula count is reported)

Each detected table goes to DuckDB (exact SQL answers) AND to text chunks
(summary chunk + row-group chunks with the header repeated) for semantic search.
"""
from __future__ import annotations

import json
import math
import uuid
from pathlib import Path
from typing import Any, Callable

import pandas as pd

from ..config import ROWS_PER_CHUNK
from ..store import Chunk, store
from ..tables import clean_columns, coerce_types, describe_df, sanitize

MAX_TEXT_ROWS = 3000   # rows per table indexed as text; SQL always covers everything
PREVIEW_ROWS, PREVIEW_COLS = 300, 60


def _empty(v: Any) -> bool:
    return v is None or (isinstance(v, float) and math.isnan(v)) or (isinstance(v, str) and not v.strip())


def _fmt(v: Any) -> str:
    if _empty(v):
        return ""
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f"{v:.4f}".rstrip("0").rstrip(".")
    if hasattr(v, "isoformat"):
        s = v.isoformat()
        return s[:10] if s.endswith("T00:00:00") else s
    return str(v).replace("\n", " ").replace("|", "/").strip()


# ---------------------------------------------------------------- grids ----
def _grids(path: Path) -> tuple[dict[str, list[list[Any]]], dict[str, dict[str, int]]]:
    ext = path.suffix.lower()
    stats: dict[str, dict[str, int]] = {}
    if ext in {".csv", ".tsv", ".txt"}:
        sep = "\t" if ext == ".tsv" else None
        df = pd.read_csv(path, header=None, dtype=object, sep=sep, engine="python", on_bad_lines="skip")
        return {path.stem: df.where(df.notna(), None).values.tolist()}, {}
    if ext in {".xlsx", ".xlsm"}:
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True)
        formulas: dict[str, int] = {}
        if path.stat().st_size < 15_000_000:
            try:
                wbf = openpyxl.load_workbook(path, data_only=False, read_only=True)
                for ws in wbf.worksheets:
                    formulas[ws.title] = sum(
                        1 for row in ws.iter_rows(values_only=True) for v in row
                        if isinstance(v, str) and v.startswith("=")
                    )
                wbf.close()
            except Exception:
                pass
        grids = {}
        for ws in wb.worksheets:
            if ws.sheet_state != "visible" and ws.max_row <= 1:
                continue
            grid = [list(r) for r in ws.iter_rows(values_only=True)]
            merged = list(ws.merged_cells.ranges)
            for mr in merged:  # propagate merged value to every cell of the range
                top = grid[mr.min_row - 1][mr.min_col - 1] if mr.min_row - 1 < len(grid) else None
                for r in range(mr.min_row - 1, min(mr.max_row, len(grid))):
                    for c in range(mr.min_col - 1, min(mr.max_col, len(grid[r]))):
                        grid[r][c] = top
            grids[ws.title] = grid
            stats[ws.title] = {"merged": len(merged), "formulas": formulas.get(ws.title, 0),
                               "hidden": int(ws.sheet_state != "visible")}
        return grids, stats
    # .xls / .ods / anything pandas can read
    sheets = pd.read_excel(path, sheet_name=None, header=None, dtype=object)
    return {name: df.where(df.notna(), None).values.tolist() for name, df in sheets.items()}, {}


# -------------------------------------------------------------- regions ----
def _regions(grid: list[list[Any]]) -> list[list[list[Any]]]:
    """Split a sheet into rectangular blocks separated by blank rows / blank columns."""
    width = max((len(r) for r in grid), default=0)
    grid = [list(r) + [None] * (width - len(r)) for r in grid]
    row_blocks, cur = [], []
    for r in grid:
        if all(_empty(v) for v in r):
            if cur:
                row_blocks.append(cur)
                cur = []
        else:
            cur.append(r)
    if cur:
        row_blocks.append(cur)

    regions = []
    for block in row_blocks:
        used = [any(not _empty(row[c]) for row in block) for c in range(width)]
        c = 0
        while c < width:
            if not used[c]:
                c += 1
                continue
            start = c
            while c < width and used[c]:
                c += 1
            regions.append([row[start:c] for row in block])
    return regions


def _is_text(v: Any) -> bool:
    return isinstance(v, str) and not _looks_numeric(v)


def _looks_numeric(s: str) -> bool:
    t = s.strip().replace(",", "").replace("$", "").replace("%", "").replace("(", "-").replace(")", "")
    try:
        float(t)
        return True
    except ValueError:
        return False


def _header_score(row: list[Any]) -> float:
    filled = [v for v in row if not _empty(v)]
    if not filled:
        return 0
    return (sum(_is_text(v) for v in filled) / len(filled)) * (len(filled) / len(row))


def _parse_region(rows: list[list[Any]]) -> dict[str, Any]:
    """Return {'kind': 'table'|'kv'|'text', ...}."""
    width = len(rows[0])
    filled_counts = [sum(not _empty(v) for v in r) for r in rows]

    # title / caption rows: leading rows with a single filled cell on a wide block
    title_lines = []
    while rows and width >= 3 and filled_counts[0] <= 1 and len(rows) > 2:
        title_lines.append(" ".join(_fmt(v) for v in rows[0] if not _empty(v)))
        rows, filled_counts = rows[1:], filled_counts[1:]

    if len(rows) == 1 or width == 1:
        text = "\n".join(" | ".join(_fmt(v) for v in r if not _empty(v)) for r in rows)
        return {"kind": "text", "text": "\n".join(title_lines + [text]).strip()}

    # form-style label/value block
    if width <= 4 and len(rows) <= 60:
        labels = [r[0] for r in rows if not _empty(r[0])]
        colon = sum(isinstance(v, str) and v.strip().endswith(":") for v in labels)
        value_types = {type(r[1]).__name__ for r in rows if len(r) > 1 and not _empty(r[1])}
        if labels and sum(_is_text(v) for v in labels) / len(labels) > 0.8 and (
                colon >= max(1, len(labels) // 2) or (width == 2 and len(rows) <= 30 and len(value_types) > 1)):
            lines = []
            for r in rows:
                vals = [_fmt(v) for v in r if not _empty(v)]
                if vals:
                    lines.append(f"{vals[0].rstrip(':')}: {' | '.join(vals[1:])}" if len(vals) > 1 else vals[0])
            return {"kind": "kv", "text": "\n".join(title_lines + lines), "title": " ".join(title_lines)}

    # header detection (search first 10 rows), with multi-row header support
    best, best_score = 0, -1.0
    for i in range(min(10, len(rows) - 1)):
        s = _header_score(rows[i])
        if s > best_score + 0.05:
            best, best_score = i, s
        if s >= 0.85:
            break
    header_rows = [rows[best]] if best_score >= 0.5 else []
    data_start = best + 1 if header_rows else 0
    # group header directly above (e.g. merged "H1" spanning "Q1 Q2")
    if header_rows and best > 0 and _header_score(rows[best - 1]) >= 0.3 \
            and sum(not _empty(v) for v in rows[best - 1]) >= 2:
        header_rows.insert(0, rows[best - 1])
        best -= 1
    if header_rows and data_start < len(rows) - 1 and _header_score(rows[data_start]) >= 0.85 \
            and _header_score(rows[data_start + 1]) < 0.6:
        header_rows.append(rows[data_start])
        data_start += 1
    pre = rows[:best] if header_rows else []
    title_lines += [" ".join(_fmt(v) for v in r if not _empty(v)) for r in pre]

    if header_rows:
        filled_rows = []
        for i, hr in enumerate(header_rows):
            hr, last = list(hr), None
            merged_already = any(not _empty(a) and a == b for a, b in zip(hr, hr[1:]))
            if i < len(header_rows) - 1 and not merged_already:  # forward-fill spanning group labels
                for j, v in enumerate(hr):
                    if _empty(v):
                        hr[j] = last
                    else:
                        last = v
            filled_rows.append(hr)
        cols = []
        for parts in zip(*filled_rows):
            labels = []
            for x in (_fmt(v) for v in parts):
                if x and x not in labels:
                    labels.append(x)
            cols.append(" / ".join(labels))
    else:
        cols = [f"col_{j + 1}" for j in range(width)]

    data = rows[data_start:]
    df = pd.DataFrame(data, columns=clean_columns(cols))
    df = df.dropna(axis=1, how="all")
    df = df.loc[~df.apply(lambda r: all(_empty(v) for v in r), axis=1)]
    df = coerce_types(df).reset_index(drop=True)
    return {"kind": "table", "df": df, "title": " — ".join(t for t in title_lines if t)}


def _md_table(df: pd.DataFrame, start: int, end: int) -> str:
    cols = [str(c) for c in df.columns]
    lines = ["| " + " | ".join(cols) + " |", "|" + "---|" * len(cols)]
    for row in df.iloc[start:end].itertuples(index=False):
        lines.append("| " + " | ".join(_fmt(v) for v in row) + " |")
    return "\n".join(lines)


def _summary_text(sheet: str, name: str, title: str, df: pd.DataFrame, cols: list[dict]) -> str:
    parts = [f"Table `{name}` on sheet '{sheet}'" + (f" — {title}" if title else ""),
             f"{len(df)} rows × {len(df.columns)} columns. Columns:"]
    for c in cols:
        if "sum" in c:
            parts.append(f"- {c['name']} (numeric): min {c['min']}, max {c['max']}, mean {c['mean']}, total {c['sum']}")
        elif "start" in c:
            parts.append(f"- {c['name']} (date): {c['start']} → {c['end']}")
        elif "top" in c:
            parts.append(f"- {c['name']} (text, {c['distinct']} distinct): e.g. {', '.join(c['top'][:4])}")
        else:
            parts.append(f"- {c['name']} ({c['type']})")
    return "\n".join(parts)


# ----------------------------------------------------------------- main ----
def ingest_excel(doc_id: str, path: Path, progress: Callable[[str, float], None]):
    progress("Reading workbook", 0.05)
    grids, stats = _grids(path)
    chunks: list[Chunk] = []
    parents: dict[str, str] = {}
    tables: list[dict[str, Any]] = []
    sheets_meta: list[dict[str, Any]] = []
    taken: set[str] = set()
    preview_dir = store.dir(doc_id) / "sheets"
    preview_dir.mkdir(exist_ok=True)

    for s_i, (sheet, grid) in enumerate(grids.items()):
        progress(f"Analysing sheet '{sheet}'", 0.1 + 0.5 * s_i / max(1, len(grids)))
        width = max((len(r) for r in grid), default=0)
        # raw grid preview for the UI (exactly what the user sees in Excel)
        prev = [[_fmt(v) for v in (list(r) + [None] * (width - len(r)))[:PREVIEW_COLS]] for r in grid[:PREVIEW_ROWS]]
        (preview_dir / f"{s_i}.json").write_text(json.dumps({"name": sheet, "rows": prev,
                                                             "total_rows": len(grid), "total_cols": width}))
        sheet_tables, notes = [], []
        for region in _regions(grid):
            parsed = _parse_region(region)
            if parsed["kind"] in {"text", "kv"}:
                if parsed["text"].strip():
                    notes.append(parsed["text"])
                continue
            df = parsed["df"]
            if df.empty or df.shape[1] == 0:
                continue
            name = sanitize(f"{sheet}" if not sheet_tables else f"{sheet}_t{len(sheet_tables) + 1}", taken)
            cols = describe_df(df)
            tables.append({"name": name, "df": df, "source": f"sheet '{sheet}'", "title": parsed["title"] or sheet,
                           "sheet": sheet})
            sheet_tables.append(name)

            summary = _summary_text(sheet, name, parsed["title"], df, cols)
            parent_id = f"{doc_id}:tbl:{name}"
            parents[parent_id] = (summary + "\n\n" + _md_table(df, 0, 40))[:6000]
            chunks.append(Chunk(id=uuid.uuid4().hex[:16], doc_id=doc_id, text=summary, kind="sheet_summary",
                                sheet=sheet, section=parsed["title"] or name, parent_id=parent_id))
            n_text = min(len(df), MAX_TEXT_ROWS)
            for start in range(0, n_text, ROWS_PER_CHUNK):
                end = min(start + ROWS_PER_CHUNK, n_text)
                chunks.append(Chunk(
                    id=uuid.uuid4().hex[:16], doc_id=doc_id, kind="rows", sheet=sheet,
                    section=parsed["title"] or name,
                    text=f"Rows {start + 1}-{end} of table `{name}` (sheet '{sheet}'):\n{_md_table(df, start, end)}",
                ))
        if notes:
            note_text = f"Sheet '{sheet}' notes / form fields:\n" + "\n".join(notes)
            pid = f"{doc_id}:notes:{s_i}"
            parents[pid] = note_text[:6000]
            for piece_start in range(0, len(note_text), 1500):
                chunks.append(Chunk(id=uuid.uuid4().hex[:16], doc_id=doc_id, kind="text", sheet=sheet,
                                    section="Notes", parent_id=pid, text=note_text[piece_start:piece_start + 1500]))
        sheets_meta.append({"index": s_i, "name": sheet, "rows": len(grid), "cols": width,
                            "tables": sheet_tables, **stats.get(sheet, {})})

    info = {"sheets": sheets_meta}
    return chunks, parents, tables, info
